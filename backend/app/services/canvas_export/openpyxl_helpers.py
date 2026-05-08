"""
openpyxl 纯工具 helper。

4.3pre.5 从 `viral_agent/services/export/export_service.py` lines 18-98 / 806-811
借鉴 4 个 helper(原作者为 viral_app 时代老同事,3 年沉淀下来的合并单元格 / MergedCell
保护套路,不绑定任何业务数据契约)。

**保留原因**: 这 4 个 helper 是纯 openpyxl 工具,不含业务逻辑,无需重写。
**下线时机**: 4.5 阶段 viral_app 彻底拆分后,本文件作为 canvas_export 的唯一归宿。
"""

from __future__ import annotations

from typing import Any

from loguru import logger
from openpyxl.cell.cell import MergedCell


def safe_set_cell_style(cell, **kwargs: Any) -> None:
    """
    安全地设置单元格样式,跳过 MergedCell。

    Args:
        cell: 单元格对象
        **kwargs: 要设置的属性(font, fill, border, alignment 等)
    """
    if isinstance(cell, MergedCell):
        return

    for attr, value in kwargs.items():
        if hasattr(cell, attr):
            setattr(cell, attr, value)


def safe_set_cell_value(ws, cell_ref: str, value: Any, **style_kwargs: Any) -> bool:
    """
    安全地设置单元格值和样式,跳过 MergedCell。

    Args:
        ws: 工作表对象
        cell_ref: 单元格引用(如 'A1', 'B2')
        value: 要设置的值
        **style_kwargs: 可选的样式属性(font, fill, border, alignment 等)

    Returns:
        bool: 是否成功设置
    """
    try:
        cell = ws[cell_ref]
        if isinstance(cell, MergedCell):
            logger.debug(f"跳过合并单元格 {cell_ref} 的写入")
            return False
        cell.value = value
        for attr, style_value in style_kwargs.items():
            if hasattr(cell, attr):
                setattr(cell, attr, style_value)
        return True
    except AttributeError as e:
        logger.debug(f"单元格 {cell_ref} 写入失败: {e}")
        return False


def safe_write_merged_cell(
    ws,
    row: int,
    col: int,
    value: Any,
    merge_cols: int = 1,
    merge_rows: int = 1,
    **style_kwargs: Any,
) -> int:
    """
    安全写入可能需要合并的单元格。

    原则:先写入左上角单元格的值,再进行合并(避免 MergedCell 只读错误)。

    Args:
        ws: 工作表对象
        row: 行号(1-based)
        col: 列号(1-based,起始列)
        value: 要写入的值
        merge_cols: 要合并的列数(默认 1 表示不跨列合并)
        merge_rows: 要合并的行数(默认 1 表示不跨行合并;4.3pre.5 新增,
                    用于 Sheet 2 矩阵里 B 列要素名的跨行合并)
        **style_kwargs: 可选的样式属性

    Returns:
        int: 实际写入的行号(方便链式调用)
    """
    try:
        cell = ws.cell(row=row, column=col)
        if not isinstance(cell, MergedCell):
            cell.value = value
            for attr, style_value in style_kwargs.items():
                if hasattr(cell, attr):
                    setattr(cell, attr, style_value)

        if merge_cols > 1 or merge_rows > 1:
            end_col = col + merge_cols - 1
            end_row = row + merge_rows - 1
            ws.merge_cells(
                start_row=row, start_column=col, end_row=end_row, end_column=end_col
            )

        return row
    except Exception as e:
        logger.debug(f"安全写入合并单元格失败 ({row}, {col}): {e}")
        return row


def safe_int(value: Any) -> int:
    """安全转换为整数(失败返回 0)。"""
    try:
        return int(value) if value else 0
    except (ValueError, TypeError):
        return 0


__all__ = [
    "safe_set_cell_style",
    "safe_set_cell_value",
    "safe_write_merged_cell",
    "safe_int",
]
