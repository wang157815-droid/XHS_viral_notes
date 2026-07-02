"""评论分析报告 Excel 导出器单元测试。

重点:验证含 Excel 非法控制字符的小红书评论内容能被安全导出,
不触发 openpyxl IllegalCharacterError(对应评论导出 500 的回归)。
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from backend.app.services.canvas_export.comment_excel_exporter import (
    _sanitize_for_xlsx,
    build_comment_excel,
)


def _v2_output_with_illegal_chars() -> dict:
    # 评论内容里混入 openpyxl 不允许的控制字符:
    #   \x07 响铃、\x0b 垂直制表、\x1f 单元分隔、\x0c 换页、\x00 NULL、\x08 退格
    bad_content = "这款\x07吸奶器真好用\x0b静音模式很赞\x1f推荐!"
    return {
        "version": "v2",
        "sheet1_comments": [
            {
                "dimension": "产品体验",
                "content": bad_content,
                "like_count": 123,
                "author": "妈妈用户\x0c",
                "is_sub_comment": False,
                "note_title": "momcozy 吸奶器测评",
                "note_url": "https://www.xiaohongshu.com/explore/abc",
            }
        ],
        "sheet2_notes": [
            {
                "url": "https://www.xiaohongshu.com/explore/abc",
                "publish_time": "2026-06-29",
                "title": "momcozy 吸奶器测评",
                "interaction_score": 5000,
                "fetched_comment_count": 200,
                "desc": "正文含\x00非法字符\x08",
                "tags": "吸奶器/母婴",
            }
        ],
    }


def test_build_comment_excel_strips_illegal_control_chars():
    """评论含 Excel 非法控制字符时,导出应剥离这些字符且不抛异常(回归 500)。"""
    data = build_comment_excel(_v2_output_with_illegal_chars())
    assert isinstance(data, (bytes, bytearray))
    assert len(data) > 0

    wb = load_workbook(io.BytesIO(data))
    ws1 = wb["评论数据源"]
    # 第 2 行第 2 列 = 第一条评论内容,非法字符应被剥离
    assert ws1.cell(row=2, column=2).value == "这款吸奶器真好用静音模式很赞推荐!"
    # 第 2 行第 4 列 = 评论者昵称,同样应被清洗
    assert ws1.cell(row=2, column=4).value == "妈妈用户"

    ws2 = wb["笔记数据源"]
    # 第 2 行第 6 列 = 笔记具体内容
    assert ws2.cell(row=2, column=6).value == "正文含非法字符"


def test_sanitize_preserves_newline_tab_and_skips_non_str():
    """清洗应保留 \\t、\\n、\\r(wrap_text 需要 \\n 换行),且对非字符串不处理。"""
    assert _sanitize_for_xlsx("第一行\n第二行\t缩进\r\n") == "第一行\n第二行\t缩进\r\n"
    assert _sanitize_for_xlsx("去掉\x00\x07\x0b\x0c\x1f") == "去掉"
    assert _sanitize_for_xlsx(123) == 123
    assert _sanitize_for_xlsx(0) == 0
    assert _sanitize_for_xlsx(None) is None
    assert _sanitize_for_xlsx("") == ""
