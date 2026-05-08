"""
4.3pre.5 Excel 导出 E2E smoke(一次性验证脚本,运行后手工删除)。

目的:在不起完整后端的情况下,验证:
1. image_fetcher 的真实 HTTP 下载管线(用 picsum.photos 公开图替代 XHS CDN)
2. 产出的 xlsx 文件无 corrupt(openpyxl round-trip load)
3. 文件可被外部 Excel/WPS 打开(需要人工确认)

运行: python backend/tests/_manual_excel_smoke.py

验证后会打印一个磁盘路径,用户可以用 WPS/Excel 手动打开核查。
"""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

# 让脚本可以直接 python xxx 跑
REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from openpyxl import load_workbook

from backend.app.domain.canvas import build_empty_canvas
from backend.app.domain.task_context import TaskContext, TaskContextWriter
from backend.app.domain.task_status import TaskStatus
from backend.app.infrastructure.repository.task_repository import TaskRecord
from backend.app.services.canvas_export import build_excel_bytes
from backend.app.services.canvas_export.image_fetcher import CoverImageFetcher


_PUBLIC_IMAGE_URLS = [
    "https://picsum.photos/seed/1/400/400",
    "https://picsum.photos/seed/2/400/400",
    "https://picsum.photos/seed/3/400/400",
]


async def test_image_pipeline():
    print("=" * 60)
    print("  [1/3] 真实 HTTP 图片下载管线(picsum.photos)")
    print("=" * 60)
    fetcher = CoverImageFetcher(timeout=8.0)
    cache = await fetcher.fetch_all(_PUBLIC_IMAGE_URLS)
    ok = 0
    for url, stream in cache.items():
        if stream is not None:
            size = len(stream.getvalue())
            print(f"  [OK] {url} -> {size:>6} bytes(已压缩)")
            ok += 1
        else:
            print(f"  [FAIL] {url}")
    print(f"\n  成功率 {ok}/{len(_PUBLIC_IMAGE_URLS)}")
    return cache


async def build_real_xlsx(img_cache):
    print("\n" + "=" * 60)
    print("  [2/3] 构造完整 task_context + canvas → 产出真实 xlsx")
    print("=" * 60)

    ctx = TaskContext(task_id="smoke_task")
    writer = TaskContextWriter(ctx)

    # 让样本的 cover_url 命中已下载的 picsum 图片,验证图片嵌入管线走通
    url_by_note = {
        "n1": _PUBLIC_IMAGE_URLS[0],
        "n2": _PUBLIC_IMAGE_URLS[1],
        "n3": _PUBLIC_IMAGE_URLS[2],
    }

    def _make_note(nid, **extra):
        note = {
            "note_id": nid,
            "title": f"抗老精华样本 {nid}",
            "nickname": "测试达人",
            "likes": 12345,
            "collects": 2345,
            "comments": 234,
            "interaction_score": 12345 + 2345 + 234,
            "media_type": "image",
            "note_type": "图集",
            "cover_url": url_by_note.get(nid, ""),
            "url": f"https://www.xiaohongshu.com/explore/{nid}",
            "sources_hit": extra.get("sources_hit", ["category_top"]),
        }
        note.update(extra)
        return note

    writer.write(
        "crawler_output",
        {
            "source": "live",
            "sample_count": 3,
            "keywords": ["抗老精华"],
            "sources": {
                "category_top": [_make_note("n1")],
                "competitor": [
                    _make_note(
                        "n2",
                        sources_hit=["competitor"],
                        seo_top10=["抗老", "细纹", "胶原"],
                        comment_hotwords_top10=["好用", "回购"],
                    )
                ],
                "top_interaction": [_make_note("n3", published_at="2024-12-01")],
                "serp_top": [_make_note("n1")],
            },
            "all_notes": [_make_note(k) for k in url_by_note],
        },
        agent_id="smoke",
    )

    writer.write(
        "multimodal_output",
        {
            "annotations": {
                "n1": {
                    "cover_type": "纯产品图",
                    "cover_text_type": "干货/经验分享",
                    "title_type": "痛点+解决方案",
                    "opening_type": "痛点切入",
                    "product_intro_type": "直接带出",
                    "product_placement_type": "融合使用感受",
                    "pain_keywords": "暗沉,细纹",
                    "content_direction": "口播单推",
                },
                "n2": {
                    "cover_type": "达人手持产品",
                    "cover_text_type": "痛点",
                    "title_type": "数字承诺",
                    "opening_type": "好状态切入",
                    "product_intro_type": "剧情中引出",
                    "product_placement_type": "效果+质地",
                    "pain_keywords": "胶原,抗老",
                    "content_direction": "剧情",
                },
                "n3": {
                    "cover_type": "前后对比",
                    "cover_text_type": "效果",
                    "title_type": "干货/经验分享",
                    "opening_type": "干货切入",
                    "product_intro_type": "干货分享引出",
                    "product_placement_type": "产品用法/使用年龄",
                    "pain_keywords": "法令纹",
                    "content_direction": "知识科普",
                },
            }
        },
        agent_id="smoke",
    )

    writer.write(
        "viral_model_output",
        {
            "models": [
                {
                    "model_id": "M1",
                    "name": "口播单推型",
                    "description": "达人手持直接推荐",
                    "coverage": 0.5,
                    "avg_interaction": 15000,
                    "paragraph_id": "M1",
                    "elements": {
                        "A_cover": [
                            {
                                "type": "达人手持产品",
                                "ratio": 0.6,
                                "count": 6,
                                "paragraph_id": "M1-A_cover-C1",
                                "examples": [
                                    {
                                        "note_id": "n1",
                                        "title": "示例",
                                        "cover_url": _PUBLIC_IMAGE_URLS[0],
                                    }
                                ],
                            },
                            {
                                "type": "纯产品图",
                                "ratio": 0.4,
                                "count": 4,
                                "paragraph_id": "M1-A_cover-C2",
                                "examples": [],
                            },
                        ],
                        "C_title": [
                            {
                                "type": "痛点+解决方案",
                                "ratio": 1.0,
                                "count": 10,
                                "paragraph_id": "M1-C_title-C1",
                                "examples": [
                                    {
                                        "note_id": "n2",
                                        "title": "示例",
                                        "cover_url": _PUBLIC_IMAGE_URLS[1],
                                    }
                                ],
                            }
                        ],
                    },
                }
            ],
            "unused_directions": [
                {
                    "direction": "日常 vlog",
                    "ratio": 0.05,
                    "avg_interaction": 20000,
                    "reason": "闭环验证差",
                }
            ],
            "total_sample_count": 20,
            "taxonomy_version": "v1",
        },
        agent_id="smoke",
    )

    writer.write(
        "semantic_output",
        {
            "content_direction": {
                "top_direction": "口播单推",
                "summary_points": ["方向集中"],
                "highlight": "口播主导",
            },
            "pain_points_top": [
                {"keyword": "暗沉", "count": 12},
                {"keyword": "细纹", "count": 8},
                {"keyword": "法令纹", "count": 5},
            ],
            "seo_aggregation": {
                "core_keywords": [{"keyword": "抗老", "count": 15}],
                "long_tail": [],
                "differentiation_advice": "聚焦细纹场景,避开 30 天承诺",
            },
            "stats_axis_label": "皮肤问题",  # 用自定义轴名验证参数化
        },
        agent_id="smoke",
    )

    record = TaskRecord(
        task_id="smoke_task",
        owner_user_id="smoke_user",
        status=TaskStatus.COMPLETED,
        keywords=["抗老精华"],
    )
    canvas = build_empty_canvas(task_id="smoke_task", title="爆文洞察 · 抗老精华(smoke)")

    data = await build_excel_bytes(
        task_record=record,
        canvas=canvas,
        task_context=ctx,
        total_timeout=60.0,
    )
    print(f"  xlsx 字节数: {len(data)}")
    return data


def verify_xlsx(data: bytes):
    print("\n" + "=" * 60)
    print("  [3/3] openpyxl round-trip 验证 + 写磁盘")
    print("=" * 60)
    wb = load_workbook(io.BytesIO(data))
    print(f"  sheet 名列表: {wb.sheetnames}")
    sheet5_name = "【品类】抗老精华互动top"
    expected = [
        "爆文总结",
        "爆文总结详情2",
        "数据源总",
        "竞品爆文",
        sheet5_name,
        "【精华】小红书前10屏爆文",
        "草稿",
    ]
    assert wb.sheetnames == expected, f"sheet 名不匹配: {wb.sheetnames}"

    # 校验 Sheet 1 H1 轴名参数化
    ws1 = wb["爆文总结"]
    assert ws1["H1"].value == "皮肤问题", f"H1 实际值: {ws1['H1'].value}"
    print(f"  [OK] Sheet 1 H1 轴名 = '皮肤问题'(参数化生效)")

    # 校验 Sheet 3 的 N-S 列映射
    ws3 = wb["数据源总"]
    assert ws3["N2"].value == "纯产品图"
    assert ws3["J2"].value == "口播单推"
    print(f"  [OK] Sheet 3 N-S 列映射正确")

    # 校验 Sheet 4 多了 T/U 列
    ws4 = wb["竞品爆文"]
    assert ws4["T1"].value == "笔记涵盖热搜词 Top10"
    print(f"  [OK] Sheet 4 额外 SEO/热词列存在")

    # 校验 Sheet 5 published_at
    ws5 = wb[sheet5_name]
    assert ws5["B2"].value == "2024-12-01"
    print(f"  [OK] Sheet 5 published_at 在 B 列")

    # 保存到磁盘
    out = REPO_ROOT / "backend" / "tests" / "_manual_excel_smoke_out.xlsx"
    out.write_bytes(data)
    print(f"\n  xlsx 保存到: {out}")
    print(f"  大小: {len(data):>6} bytes")
    print(f"  请用 WPS / Excel 手动打开查看,确认不提示 corrupt。")


async def main():
    img_cache = await test_image_pipeline()
    data = await build_real_xlsx(img_cache)
    verify_xlsx(data)
    print("\n" + "=" * 60)
    print("  4.3pre.5 E2E smoke 完成")
    print("=" * 60)


if __name__ == "__main__":
    asyncio.run(main())
