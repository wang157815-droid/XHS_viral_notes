"""
多模式导出服务
支持图文/视频分开导出为独立Excel文件
"""
import json
import os
from datetime import datetime
from typing import Dict, Any, List, Union
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from loguru import logger


def export_with_mode(
    analysis_file_path: str,
    export_mode: str = "combined"
) -> Union[str, Dict[str, str]]:
    """
    根据导出模式导出分析结果

    Args:
        analysis_file_path: 分析结果JSON文件路径
        export_mode: 导出模式
            - "combined": 单文件导出（默认，现有行为）
            - "separate": 图文+视频两个独立Excel
            - "image_only": 仅导出图文报告
            - "video_only": 仅导出视频报告

    Returns:
        - combined/image_only/video_only: 返回单个文件路径 (str)
        - separate: 返回包含多个文件路径的字典 (Dict[str, str])
    """
    # 读取分析数据
    with open(analysis_file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    output_dir = os.path.dirname(analysis_file_path)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    if export_mode == "separate":
        return _export_separate_files(data, output_dir, timestamp)
    elif export_mode == "image_only":
        return _export_single_type(data, output_dir, timestamp, "image")
    elif export_mode == "video_only":
        return _export_single_type(data, output_dir, timestamp, "video")
    else:
        # combined 模式：使用原有导出逻辑
        from viral_agent.services.export.export_service import export_to_excel
        return export_to_excel(analysis_file_path)


def _export_separate_files(
    data: Dict[str, Any],
    output_dir: str,
    timestamp: str
) -> Dict[str, str]:
    """
    导出为两个独立的Excel文件（图文+视频）

    Returns:
        包含文件路径的字典：
        {
            'image': '图文报告路径',
            'video': '视频报告路径',
            'comparison': '对比报告路径'（可选）
        }
    """
    result = {}

    # 获取类型分离数据
    image_features = data.get('image_note_features', {})
    video_features = data.get('video_note_features', {})
    type_summary = data.get('type_summary', {})

    image_count = image_features.get('count', 0)
    video_count = video_features.get('count', 0)

    # 1. 导出图文报告
    if image_count > 0:
        image_path = _export_single_type(data, output_dir, timestamp, "image")
        result['image'] = image_path
        logger.info(f"✅ 图文报告已生成: {image_path}")
    else:
        logger.warning("⚠️ 没有图文笔记，跳过图文报告")

    # 2. 导出视频报告
    if video_count > 0:
        video_path = _export_single_type(data, output_dir, timestamp, "video")
        result['video'] = video_path
        logger.info(f"✅ 视频报告已生成: {video_path}")
    else:
        logger.warning("⚠️ 没有视频笔记，跳过视频报告")

    # 3. 生成对比报告（如果两种类型都有数据）
    if image_count > 0 and video_count > 0:
        comparison_path = _create_comparison_report(
            data, output_dir, timestamp, type_summary
        )
        result['comparison'] = comparison_path
        logger.info(f"✅ 对比报告已生成: {comparison_path}")

    return result


def _export_single_type(
    data: Dict[str, Any],
    output_dir: str,
    timestamp: str,
    note_type: str
) -> str:
    """
    导出单一类型的Excel报告

    Args:
        data: 完整分析数据
        output_dir: 输出目录
        timestamp: 时间戳
        note_type: "image" 或 "video"

    Returns:
        生成的Excel文件路径
    """
    # 获取类型专属特征
    if note_type == "image":
        type_features = data.get('image_note_features', {})
        type_label = "图文"
        filename = f"viral_report_image_{timestamp}.xlsx"
    else:
        type_features = data.get('video_note_features', {})
        type_label = "视频"
        filename = f"viral_report_video_{timestamp}.xlsx"

    # 创建工作簿
    wb = Workbook()

    # 创建类型专属报告工作表
    _create_type_overview_sheet(wb, data, type_features, type_label, note_type)
    _create_type_top_notes_sheet(wb, type_features, type_label)

    # 如果是图文，添加图文专属分析
    if note_type == "image":
        _add_image_specific_sheets(wb, data)
    # 如果是视频，添加视频专属分析
    else:
        _add_video_specific_sheets(wb, data)

    # 删除默认空Sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    # 保存文件
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)

    return output_path


def _create_type_overview_sheet(
    wb: Workbook,
    data: Dict[str, Any],
    type_features: Dict[str, Any],
    type_label: str,
    note_type: str
):
    """创建类型专属总览工作表"""
    ws = wb.create_sheet(f"📊 {type_label}分析总览")

    # 样式定义
    header_font = Font(bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E86AB" if note_type == "image" else "A23B72")
    normal_font = Font(size=11)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    row = 1

    # 标题
    ws.merge_cells(f'A{row}:D{row}')
    ws[f'A{row}'] = f"📈 {type_label}爆款笔记分析报告"
    ws[f'A{row}'].font = Font(bold=True, size=16)
    row += 2

    # 基础统计
    ws[f'A{row}'] = "📊 基础统计"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    stats_items = [
        ("笔记数量", f"{type_features.get('count', 0)} 篇"),
        ("关键词", data.get('keyword', 'N/A')),
        ("分析时间", data.get('analysis_time', 'N/A')),
    ]

    for label, value in stats_items:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = normal_font
        ws[f'B{row}'].font = normal_font
        ws[f'A{row}'].border = thin_border
        ws[f'B{row}'].border = thin_border
        row += 1

    row += 1

    # 互动统计
    ws[f'A{row}'] = "💬 互动统计"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    interaction_stats = type_features.get('interaction_stats', {})
    interaction_items = [
        ("平均互动量", f"{interaction_stats.get('avg', 0):,.0f}"),
        ("最高互动量", f"{interaction_stats.get('max', 0):,}"),
        ("最低互动量", f"{interaction_stats.get('min', 0):,}"),
        ("总互动量", f"{interaction_stats.get('total', 0):,}"),
    ]

    for label, value in interaction_items:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].border = thin_border
        ws[f'B{row}'].border = thin_border
        row += 1

    row += 1

    # 标题特征
    ws[f'A{row}'] = "📝 标题特征"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    title_features = type_features.get('title_features', {})
    ws[f'A{row}'] = "平均标题长度"
    ws[f'B{row}'] = f"{title_features.get('avg_length', 0):.1f} 字"
    ws[f'A{row}'].border = thin_border
    ws[f'B{row}'].border = thin_border
    row += 1

    top_keywords = title_features.get('top_keywords', [])
    ws[f'A{row}'] = "高频关键词"
    ws[f'B{row}'] = ", ".join(top_keywords[:5]) if top_keywords else "无"
    ws[f'A{row}'].border = thin_border
    ws[f'B{row}'].border = thin_border
    row += 1

    row += 1

    # 发布时间
    ws[f'A{row}'] = "⏰ 发布时间分析"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    time_features = type_features.get('time_features', {})
    ws[f'A{row}'] = "最佳发布时间"
    ws[f'B{row}'] = time_features.get('best_publish_time', '未知')
    ws[f'A{row}'].border = thin_border
    ws[f'B{row}'].border = thin_border

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 40


def _create_type_top_notes_sheet(
    wb: Workbook,
    type_features: Dict[str, Any],
    type_label: str
):
    """创建TOP笔记工作表"""
    ws = wb.create_sheet(f"🏆 TOP{type_label}笔记")

    # 表头
    headers = ["排名", "笔记ID", "标题", "点赞", "收藏", "评论", "总互动"]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # 数据
    top_notes = type_features.get('top_notes', [])
    for idx, note in enumerate(top_notes, 1):
        row = idx + 1
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=note.get('note_id', ''))
        ws.cell(row=row, column=3, value=note.get('title', ''))
        ws.cell(row=row, column=4, value=note.get('liked_count', 0))
        ws.cell(row=row, column=5, value=note.get('collected_count', 0))
        ws.cell(row=row, column=6, value=note.get('comment_count', 0))
        ws.cell(row=row, column=7, value=note.get('interaction_score', 0))

    # 设置列宽
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12


def _add_image_specific_sheets(wb: Workbook, data: Dict[str, Any]):
    """添加图文专属分析工作表"""
    # 封面分析
    cover_features = data.get('cover_features', {})
    if cover_features and cover_features.get('status') != 'no_data':
        ws = wb.create_sheet("🖼️ 封面分析")
        _write_dict_to_sheet(ws, cover_features, "封面特征分析")

    # OCR分析
    all_images_ocr = data.get('all_images_ocr', {})
    if all_images_ocr and all_images_ocr.get('status') != 'no_data':
        ws = wb.create_sheet("📸 配图OCR分析")
        _write_dict_to_sheet(ws, all_images_ocr, "配图文字识别结果")


def _add_video_specific_sheets(wb: Workbook, data: Dict[str, Any]):
    """添加视频专属分析工作表"""
    viral_model = data.get('viral_model', {})
    video_ai_insights = viral_model.get('video_ai_insights', {})

    if video_ai_insights and video_ai_insights.get('status') not in ['no_videos', 'skipped', 'error']:
        ws = wb.create_sheet("🎬 视频AI分析")
        _write_dict_to_sheet(ws, video_ai_insights, "视频AI深度分析结果")


def _create_comparison_report(
    data: Dict[str, Any],
    output_dir: str,
    timestamp: str,
    type_summary: Dict[str, Any]
) -> str:
    """创建图文vs视频对比报告"""
    wb = Workbook()
    ws = wb.active
    ws.title = "📊 图文vs视频对比"

    # 样式
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    row = 1

    # 标题
    ws.merge_cells('A1:D1')
    ws['A1'] = "📈 图文 vs 视频 对比分析报告"
    ws['A1'].font = Font(bold=True, size=16)
    row = 3

    # 数量对比
    ws[f'A{row}'] = "📊 数量分布"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    # 表头
    for col, header in enumerate(['指标', '图文', '视频', '总计'], 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    row += 1

    # 数量数据
    ws.cell(row=row, column=1, value="笔记数量").border = thin_border
    ws.cell(row=row, column=2, value=type_summary.get('image_count', 0)).border = thin_border
    ws.cell(row=row, column=3, value=type_summary.get('video_count', 0)).border = thin_border
    ws.cell(row=row, column=4, value=type_summary.get('total_count', 0)).border = thin_border
    row += 1

    ws.cell(row=row, column=1, value="占比").border = thin_border
    ws.cell(row=row, column=2, value=f"{type_summary.get('image_percentage', 0)}%").border = thin_border
    ws.cell(row=row, column=3, value=f"{type_summary.get('video_percentage', 0)}%").border = thin_border
    ws.cell(row=row, column=4, value="100%").border = thin_border
    row += 2

    # 互动对比
    ws[f'A{row}'] = "💬 互动数据对比"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    for col, header in enumerate(['指标', '图文', '视频', '差异'], 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    row += 1

    image_stats = type_summary.get('image_stats', {})
    video_stats = type_summary.get('video_stats', {})
    comparison = type_summary.get('comparison', {})

    ws.cell(row=row, column=1, value="平均互动量").border = thin_border
    ws.cell(row=row, column=2, value=f"{image_stats.get('avg_interaction', 0):,.0f}").border = thin_border
    ws.cell(row=row, column=3, value=f"{video_stats.get('avg_interaction', 0):,.0f}").border = thin_border

    higher = comparison.get('higher_interaction', 'unknown')
    diff = comparison.get('interaction_diff_percent', 0)
    if higher == 'video':
        diff_text = f"视频高 {diff}%"
    elif higher == 'image':
        diff_text = f"图文高 {diff}%"
    else:
        diff_text = "相当"
    ws.cell(row=row, column=4, value=diff_text).border = thin_border
    row += 1

    ws.cell(row=row, column=1, value="最高互动量").border = thin_border
    ws.cell(row=row, column=2, value=f"{image_stats.get('max_interaction', 0):,}").border = thin_border
    ws.cell(row=row, column=3, value=f"{video_stats.get('max_interaction', 0):,}").border = thin_border
    ws.cell(row=row, column=4, value="-").border = thin_border
    row += 2

    # 结论与建议
    ws[f'A{row}'] = "💡 结论与建议"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1

    ws.merge_cells(f'A{row}:D{row}')
    recommendation = comparison.get('recommendation', '暂无建议')
    ws[f'A{row}'] = recommendation
    ws[f'A{row}'].font = Font(size=11, italic=True)
    ws[f'A{row}'].alignment = Alignment(wrap_text=True)

    # 列宽
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 20

    # 保存
    filename = f"viral_report_comparison_{timestamp}.xlsx"
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)

    return output_path


def _write_dict_to_sheet(ws, data: Dict[str, Any], title: str, start_row: int = 1):
    """将字典数据写入工作表（简化版）"""
    ws.cell(row=start_row, column=1, value=title)
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=14)

    row = start_row + 2
    for key, value in data.items():
        if isinstance(value, dict):
            ws.cell(row=row, column=1, value=f"【{key}】")
            ws.cell(row=row, column=1).font = Font(bold=True)
            row += 1
            for sub_key, sub_value in value.items():
                ws.cell(row=row, column=1, value=f"  {sub_key}")
                if isinstance(sub_value, (list, dict)):
                    ws.cell(row=row, column=2, value=str(sub_value)[:500])
                else:
                    ws.cell(row=row, column=2, value=sub_value)
                row += 1
        elif isinstance(value, list):
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=str(value)[:500])
            row += 1
        else:
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)
            row += 1

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 80
