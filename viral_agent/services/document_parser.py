"""
文档解析服务
支持多格式文档解析：PDF、Word、Markdown、TXT
"""
from typing import Dict, Any, List
from pathlib import Path
import re
from loguru import logger


class DocumentParser:
    """文档解析器"""

    def __init__(self):
        """初始化解析器"""
        self.supported_formats = ['.pdf', '.docx', '.doc', '.md', '.txt']

    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """
        解析文档，提取文本内容和元数据

        Args:
            file_path: 文件路径

        Returns:
            {
                'text': '文档全文',
                'metadata': {
                    'title': '文档标题',
                    'pages': 10,
                    'format': 'pdf',
                    'word_count': 1000
                }
            }

        Raises:
            FileNotFoundError: 文件不存在
            ValueError: 不支持的文件格式
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()

        if ext not in self.supported_formats:
            raise ValueError(f"不支持的文件格式: {ext}。支持的格式: {', '.join(self.supported_formats)}")

        logger.info(f"开始解析文档: {path.name} (格式: {ext})")

        try:
            if ext == '.pdf':
                result = self._parse_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                result = self._parse_word(file_path)
            elif ext == '.md':
                result = self._parse_markdown(file_path)
            elif ext == '.txt':
                result = self._parse_txt(file_path)
            else:
                raise ValueError(f"未实现的格式: {ext}")

            # 添加通用元数据
            result['metadata']['filename'] = path.name
            result['metadata']['format'] = ext[1:]  # 去掉点号
            result['metadata']['word_count'] = len(result['text'])

            logger.success(f"文档解析完成: {path.name}, 字数: {result['metadata']['word_count']}")
            return result

        except Exception as e:
            logger.error(f"文档解析失败 {path.name}: {e}")
            raise

    def _parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        解析PDF文档

        Args:
            file_path: PDF文件路径

        Returns:
            解析结果
        """
        try:
            import pypdf
        except ImportError:
            raise ImportError("请先安装pypdf: pip install pypdf")

        text_parts = []
        metadata = {}

        try:
            with open(file_path, 'rb') as f:
                pdf_reader = pypdf.PdfReader(f)

                # 提取元数据
                metadata['pages'] = len(pdf_reader.pages)
                if pdf_reader.metadata:
                    metadata['title'] = pdf_reader.metadata.get('/Title', '')
                    metadata['author'] = pdf_reader.metadata.get('/Author', '')

                # 提取文本
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

            full_text = '\n\n'.join(text_parts)

            return {
                'text': full_text,
                'metadata': metadata
            }

        except Exception as e:
            logger.error(f"PDF解析失败: {e}")
            raise

    def _parse_word(self, file_path: str) -> Dict[str, Any]:
        """
        解析Word文档

        Args:
            file_path: Word文件路径

        Returns:
            解析结果
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError("请先安装python-docx: pip install python-docx")

        try:
            doc = Document(file_path)

            # 提取段落文本
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # 提取表格文本
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)

            full_text = '\n\n'.join(text_parts)

            # 提取元数据
            metadata = {
                'paragraphs': len(doc.paragraphs),
                'tables': len(doc.tables)
            }

            # 尝试提取文档属性
            if hasattr(doc.core_properties, 'title'):
                metadata['title'] = doc.core_properties.title or ''
            if hasattr(doc.core_properties, 'author'):
                metadata['author'] = doc.core_properties.author or ''

            return {
                'text': full_text,
                'metadata': metadata
            }

        except Exception as e:
            logger.error(f"Word文档解析失败: {e}")
            raise

    def _parse_markdown(self, file_path: str) -> Dict[str, Any]:
        """
        解析Markdown文档

        Args:
            file_path: Markdown文件路径

        Returns:
            解析结果
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 提取标题（第一个一级标题作为title）
            title_match = re.search(r'^#\s+(.+)$', content, re.MULTILINE)
            title = title_match.group(1) if title_match else ''

            # 统计标题数量
            h1_count = len(re.findall(r'^#\s+', content, re.MULTILINE))
            h2_count = len(re.findall(r'^##\s+', content, re.MULTILINE))

            metadata = {
                'title': title,
                'h1_count': h1_count,
                'h2_count': h2_count,
                'has_code_blocks': '```' in content
            }

            return {
                'text': content,
                'metadata': metadata
            }

        except Exception as e:
            logger.error(f"Markdown解析失败: {e}")
            raise

    def _parse_txt(self, file_path: str) -> Dict[str, Any]:
        """
        解析纯文本文档

        Args:
            file_path: TXT文件路径

        Returns:
            解析结果
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 简单元数据
            lines = content.split('\n')
            metadata = {
                'lines': len(lines),
                'title': lines[0][:50] if lines else ''  # 第一行前50字作为标题
            }

            return {
                'text': content,
                'metadata': metadata
            }

        except UnicodeDecodeError:
            # 尝试其他编码
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    content = f.read()

                lines = content.split('\n')
                metadata = {
                    'lines': len(lines),
                    'title': lines[0][:50] if lines else '',
                    'encoding': 'gbk'
                }

                return {
                    'text': content,
                    'metadata': metadata
                }
            except Exception as e:
                logger.error(f"TXT文件编码错误: {e}")
                raise

    def split_text(
        self,
        text: str,
        chunk_size: int = 500,
        overlap: int = 50,
        separators: List[str] = None
    ) -> List[str]:
        """
        文本分块，用于向量化

        Args:
            text: 原始文本
            chunk_size: 分块大小（字符数）
            overlap: 重叠字符数（保持上下文连贯）
            separators: 分隔符列表（按优先级）

        Returns:
            文本块列表
        """
        if separators is None:
            separators = ['\n\n', '\n', '。', '！', '？', '. ', '! ', '? ']

        chunks = []
        separator = ""  # 初始化分隔符

        # 首先按优先级最高的分隔符分割
        for sep in separators:
            if sep in text:
                separator = sep
                parts = text.split(separator)
                break
        else:
            # 如果没有找到分隔符，直接使用整个文本
            parts = [text]

        current_chunk = ""

        for part in parts:
            # 如果当前块加上新部分不超过chunk_size，就合并
            if len(current_chunk) + len(part) + len(separator) <= chunk_size:
                current_chunk += part + separator
            else:
                # 否则，保存当前块，开始新块
                if current_chunk:
                    chunks.append(current_chunk.strip())

                # 新块从overlap位置开始
                if overlap > 0 and current_chunk:
                    overlap_text = current_chunk[-overlap:]
                    current_chunk = overlap_text + part + separator
                else:
                    current_chunk = part + separator

        # 添加最后一块
        if current_chunk:
            chunks.append(current_chunk.strip())

        # 过滤空块
        chunks = [c for c in chunks if c.strip()]

        logger.debug(f"文本分块完成: {len(chunks)} 个块")
        return chunks

    def extract_keywords(self, text: str, top_k: int = 10) -> List[str]:
        """
        提取文本关键词（简单实现）

        Args:
            text: 文本内容
            top_k: 返回前K个关键词

        Returns:
            关键词列表
        """
        try:
            import jieba
            import jieba.analyse
        except ImportError:
            logger.warning("jieba未安装，跳过关键词提取")
            return []

        try:
            keywords = jieba.analyse.extract_tags(text, topK=top_k, withWeight=False)
            return keywords
        except Exception as e:
            logger.error(f"关键词提取失败: {e}")
            return []
