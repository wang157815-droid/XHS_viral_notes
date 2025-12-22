# -*- coding: utf-8 -*-
"""读取PRD Excel文件并保存为UTF-8文件"""
import openpyxl

def main():
    file_path = r"C:\Users\ASUS\Desktop\XHS_Agent.xlsx"
    output_path = r"C:\Users\ASUS\PycharmProjects\Spider_XHS\scripts\prd_content.txt"

    wb = openpyxl.load_workbook(file_path)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("=" * 60 + "\n")
        f.write(f"工作表列表: {wb.sheetnames}\n")
        f.write("=" * 60 + "\n")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            f.write(f"\n{'='*20} {sheet_name} {'='*20}\n")

            for row_idx, row in enumerate(ws.iter_rows(max_row=80, values_only=True), 1):
                if any(cell is not None for cell in row):
                    display = []
                    for cell in row:
                        if cell is None:
                            display.append("")
                        elif len(str(cell)) > 60:
                            display.append(str(cell)[:60] + "...")
                        else:
                            display.append(str(cell))
                    f.write(f"{row_idx:3}: {display}\n")

    print(f"内容已保存到: {output_path}")

if __name__ == "__main__":
    main()
